"""DOCX 采样预检：摄入 DOCX 文件路径，产出 PrecheckResult（分流决策 + 文档统计[字符数/段落数/表格数/图片数]）。

决策域: {WHOLE_TEXT_PIPELINE, SKIP_TEXT_PIPELINE}

总字符 >= 150 -> WHOLE_TEXT_PIPELINE；
总字符 < 150 且含图 -> SKIP_TEXT_PIPELINE + vlm_candidate；
总字符 < 150 且无图 -> SKIP_TEXT_PIPELINE。
图片量度用数量（python-docx 难拿面积）。
"""
from pathlib import Path

from config import DOCX_TEXT_THRESHOLD
from .result import DispatchDecision, PrecheckResult


def _count_images(document) -> int:
    """统计文档内嵌图片数：python-docx 无直接 API，数 inline_shapes。"""
    try:
        return len(document.inline_shapes)
    except Exception:
        return 0


def precheck_docx(path: Path) -> PrecheckResult:
    """对 DOCX 做预检：统计段落/表格字符与图片数并产出分流决策。

    Args:
        path: DOCX 文件路径。

    Returns:
        PrecheckResult：WHOLE_TEXT_PIPELINE 或 SKIP_TEXT_PIPELINE 决策与质量信号。
    """
    from docx import Document

    document = Document(str(path))
    total_chars = 0
    paragraph_count = 0

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            paragraph_count += 1
            total_chars += len(paragraph.text)

    table_count = len(document.tables)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                total_chars += len(cell.text)

    image_count = _count_images(document)

    if total_chars >= DOCX_TEXT_THRESHOLD:
        decision = DispatchDecision.WHOLE_TEXT_PIPELINE
        vlm_candidates = 0
    elif image_count > 0:
        decision = DispatchDecision.SKIP_TEXT_PIPELINE
        vlm_candidates = image_count
    else:
        decision = DispatchDecision.SKIP_TEXT_PIPELINE
        vlm_candidates = 0

    return PrecheckResult(
        doc_decision=decision,
        empty_page_ratio=0.0,
        vlm_candidate_count=vlm_candidates,
        degraded_flags=[],
        sampling_format_stats={
            "total_chars": total_chars,
            "paragraph_count": paragraph_count,
            "table_count": table_count,
            "image_count": image_count,
        },
    )
