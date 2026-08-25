"""测试用多格式样例生成器（008-1 precheck / 008-2 loader 共用）。

不引入额外 PDF writer：PDF 用手写最小合法文件（xref 偏移正确），
DOCX/PPTX 用 python-docx / python-pptx 生成，HTML 直接写字符串。
"""
from pathlib import Path

# 1x1 透明 PNG（用于 DOCX/PPTX 图片样例）
_TINY_PNG = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJ"
    "AAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def tiny_png_bytes() -> bytes:
    """返回 1x1 透明 PNG 字节（写临时文件后喂给 add_picture）。"""
    import base64
    return base64.b64decode(_TINY_PNG)


def make_pdf(pages: list[str] | str, content_streams: list[str] | None = None) -> bytes:
    """生成文本 PDF。pages 为每页文本；content_streams 可传自定义内容流覆盖（如多栏坐标）。"""
    if isinstance(pages, str):
        pages = [pages]
    page_count = len(pages)
    catalog, pages_obj = 1, 2
    page_num = 3
    content_num = page_num + page_count
    font_num = content_num + page_count
    total = font_num + 1
    bodies = {
        catalog: b"<< /Type /Catalog /Pages 2 0 R >>",
        pages_obj: (
            f"<< /Type /Pages /Kids [{' '.join(f'{page_num+i} 0 R' for i in range(page_count))}] "
            f"/Count {page_count} >>"
        ).encode(),
    }
    for i in range(page_count):
        page_ref = page_num + i
        content_ref = content_num + i
        bodies[page_ref] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_ref} 0 R "
            f"/Resources << /Font << /F1 {font_num} 0 R >> >> >>"
        ).encode()
        if content_streams is not None:
            stream = content_streams[i].encode()
        else:
            stream = f"BT /F1 12 Tf 72 {720 - i * 40} Td ({pages[i]}) Tj ET\n".encode()
        bodies[content_ref] = f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"endstream"
    bodies[font_num] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
    out = bytearray(b"%PDF-1.4\n")
    offsets = {}
    for num in range(1, total):
        offsets[num] = len(out)
        out.extend(f"{num} 0 obj\n".encode())
        out.extend(bodies[num])
        out.extend(b"\nendobj\n")
    xref = len(out)
    out.extend(f"xref\n0 {total}\n".encode())
    out.extend(b"0000000000 65535 f \n")
    for num in range(1, total):
        out.extend(f"{offsets[num]:010d} 00000 n \n".encode())
    out.extend(f"trailer\n<< /Size {total} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return bytes(out)


def multi_column_content_stream() -> str:
    """两栏内容流：词分布在 x=72 与 x=520 两带，触发多栏检测。"""
    left = " ".join(f"L{i}" for i in range(20))
    right = " ".join(f"R{i}" for i in range(20))
    return (
        f"BT /F1 10 Tf 72 700 Td ({left}) Tj ET\n"
        f"BT /F1 10 Tf 520 700 Td ({right}) Tj ET\n"
    )


def make_html(title: str, body: str, images: int = 0, nav: str = "") -> str:
    """生成 HTML 字符串。images 控制 <img> 数量，nav 追加样板导航。"""
    imgs = "".join("<img src='x.png'/>" for _ in range(images))
    nav_html = f"<nav>{nav}</nav>" if nav else ""
    return (
        f"<!DOCTYPE html><html><head><title>{title}</title></head>"
        f"<body>{nav_html}<article><h1>{title}</h1><p>{body}</p>{imgs}</article></body></html>"
    )


def make_docx(paragraphs: list[str], with_image: bool = False, tmp_path: Path | None = None) -> Path:
    """生成 DOCX 文件并返回路径。with_image=True 时插入一张图片。"""
    import docx

    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if with_image:
        image_path = tmp_path / "img.png"
        image_path.write_bytes(tiny_png_bytes())
        document.add_picture(str(image_path))
    out_path = tmp_path / "sample.docx"
    document.save(str(out_path))
    return out_path


def make_pptx(slides: list[list[str]], with_image_slide: bool = False,
              tmp_path: Path | None = None) -> Path:
    """生成 PPTX 文件并返回路径。每 slide 为一组文本行；with_image_slide 追加图文 slide。"""
    from pptx import Presentation

    presentation = Presentation()
    for lines in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])  # blank
        for i, line in enumerate(lines):
            box = slide.shapes.add_textbox(100000 + i * 100000, 100000, 5000000, 400000)
            box.text = line
    if with_image_slide:
        image_path = tmp_path / "img.png"
        image_path.write_bytes(tiny_png_bytes())
        image_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        image_slide.shapes.add_picture(str(image_path), 100000, 100000, 4000000, 4000000)
    out_path = tmp_path / "sample.pptx"
    presentation.save(str(out_path))
    return out_path
