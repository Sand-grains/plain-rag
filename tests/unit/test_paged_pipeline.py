"""unit: indexing/parse_backends/paged_pipeline 页级混合重试 + VLM 修正 + 跨页合并(012-2 2.4)。"""
from pathlib import Path

import pytest

from indexing.parse_backends import paged_pipeline as pp


@pytest.fixture
def auto_install_fakes():
    """覆盖 conftest 的 autouse fake 安装: 本模块只测页级管线逻辑, 无需重依赖。"""
    yield


def _quality_ok(markdown, original_table_count=None, expected_code_blocks=None):
    """fake 质量判定: 非空即达标。"""
    reasons = [] if markdown.strip() else ["empty"]
    return (not reasons, reasons, {"table_retention": None})


class TestParsePage:
    def test_chain_ok_returns_without_vlm(self):
        chain = lambda path: ("# 页\n正文", {"backend": "docling"})
        vlm = lambda img: ("VLM", {"backend": "vlm"})
        markdown, meta = pp._parse_page("p.pdf", chain, _quality_ok, vlm,
                                        lambda p, w: "img.png", "work", None, None)
        assert markdown == "# 页\n正文"
        assert meta["backend"] == "docling"
        assert "vlm_correction_applied" not in meta

    def test_chain_fails_quality_then_vlm_corrects(self):
        chain = lambda path: ("", {"backend": "docling"})  # 空产出 -> 质量不达标
        vlm = lambda img: ("# VLM 修正", {"backend": "vlm"})
        markdown, meta = pp._parse_page("p.pdf", chain, _quality_ok, vlm,
                                        lambda p, w: "img.png", "work", None, None)
        assert markdown == "# VLM 修正"
        assert meta["vlm_correction_applied"] is True

    def test_chain_raises_then_vlm_corrects(self):
        def chain(path):
            raise RuntimeError("后端失败")
        vlm = lambda img: ("# VLM", {"backend": "vlm"})
        markdown, meta = pp._parse_page("p.pdf", chain, _quality_ok, vlm,
                                        lambda p, w: "img.png", "work", None, None)
        assert markdown == "# VLM"
        assert meta["vlm_correction_applied"] is True

    def test_no_render_returns_empty(self):
        chain = lambda path: ("", {"backend": "docling"})
        vlm = lambda img: ("# VLM", {"backend": "vlm"})
        markdown, meta = pp._parse_page("p.txt", chain, _quality_ok, vlm,
                                        lambda p, w: None, "work", None, None)
        assert markdown == ""
        assert meta["vlm_correction_applied"] is False


class TestWholeDocVlmCorrect:
    def test_chain_ok_returns(self):
        chain = lambda path: ("# 整篇\n正文", {"backend": "docling"})
        markdown, meta = pp._whole_doc_vlm_correct(
            "a.html", chain, _quality_ok, lambda img: ("VLM", {}),
            lambda p, w: "img.png", "work", None, None)
        assert markdown == "# 整篇\n正文"

    def test_chain_fails_then_vlm_corrects(self):
        chain = lambda path: ("", {"backend": "docling"})
        markdown, meta = pp._whole_doc_vlm_correct(
            "a.html", chain, _quality_ok, lambda img: ("# VLM 整篇", {"backend": "vlm"}),
            lambda p, w: "img.png", "work", None, None)
        assert markdown == "# VLM 整篇"
        assert meta["vlm_correction_applied"] is True


class TestParseTextPipelinePaged:
    def test_pages_merged_with_cross_page_meta(self):
        # 两页: 第一页表格, 第二页表格续页(列数一致) -> 跨页合并
        splitter = lambda path, work: ["p0.pdf", "p1.pdf"]
        chain = lambda path: ("| a |\n|---|\n| 1 |", {"backend": "docling"})
        markdown, meta = pp.parse_text_pipeline_paged(
            "x.pdf", None, None, chain, _quality_ok, lambda img: ("VLM", {}),
            page_splitter=splitter, page_renderer=lambda p, w: None)
        assert "| 1 |" in markdown
        assert meta["page_retry_stats"]["pages"] == 2
        assert meta["cross_page_merged"] is True

    def test_vlm_corrected_pages_counted(self):
        splitter = lambda path, work: ["p0.pdf", "p1.pdf"]
        def chain(path):
            # 第一页空(质量不达标), 第二页达标
            if path == "p0.pdf":
                return ("", {"backend": "docling"})
            return ("# 页2\n正文", {"backend": "docling"})
        vlm = lambda img: ("# VLM 修正", {"backend": "vlm"})
        markdown, meta = pp.parse_text_pipeline_paged(
            "x.pdf", None, None, chain, _quality_ok, vlm,
            page_splitter=splitter, page_renderer=lambda p, w: "img.png")
        assert meta["page_retry_stats"]["vlm_corrected"] == 1
        assert "# VLM 修正" in markdown

    def test_no_page_boundary_uses_whole_doc_vlm(self):
        splitter = lambda path, work: None  # DOCX/HTML 无稳定页边界
        chain = lambda path: ("", {"backend": "docling"})
        vlm = lambda img: ("# 整篇 VLM", {"backend": "vlm"})
        markdown, meta = pp.parse_text_pipeline_paged(
            "a.html", None, None, chain, _quality_ok, vlm,
            page_splitter=splitter, page_renderer=lambda p, w: "img.png")
        assert markdown == "# 整篇 VLM"
        assert meta["vlm_correction_applied"] is True

    def test_page_chain_used_for_pages_chain_for_whole_doc(self):
        # M2: 页级用 page_chain(关闭表格/代码保留率), 整篇重试用 chain(保留整篇基准)
        splitter = lambda path, work: ["p0.pdf"]
        chain_calls, page_chain_calls = [], []
        def chain(path):
            chain_calls.append(path)
            return ("# 整篇", {"backend": "docling"})
        def page_chain(path):
            page_chain_calls.append(path)
            return ("# 页", {"backend": "docling"})
        markdown, meta = pp.parse_text_pipeline_paged(
            "x.pdf", None, None, chain, _quality_ok, lambda img: ("VLM", {}),
            page_splitter=splitter, page_renderer=lambda p, w: None, page_chain=page_chain)
        assert page_chain_calls == ["p0.pdf"]
        assert chain_calls == []

    def test_page_quality_uses_none_counts(self):
        # M2: 页级判定关闭表格/代码保留率(传 None,None), 避免整篇基准假性降级
        splitter = lambda path, work: ["p0.pdf"]
        captured = {}
        def quality_fn(markdown, original_table_count=None, expected_code_blocks=None):
            captured["table"] = original_table_count
            captured["code"] = expected_code_blocks
            return (True, [], {})
        pp.parse_text_pipeline_paged(
            "x.pdf", 5, 3, lambda p: ("# 页", {"backend": "docling"}), quality_fn,
            lambda img: ("VLM", {}), page_splitter=splitter, page_renderer=lambda p, w: None)
        assert captured == {"table": None, "code": None}

    def test_whole_doc_uses_chain_not_page_chain(self):
        # M2: 无页边界(DOCX/HTML)整篇重试走 chain, 不走 page_chain
        splitter = lambda path, work: None
        chain_calls, page_chain_calls = [], []
        def chain(path):
            chain_calls.append(path)
            return ("# 整篇", {"backend": "docling"})
        def page_chain(path):
            page_chain_calls.append(path)
            return ("# 页", {"backend": "docling"})
        markdown, meta = pp.parse_text_pipeline_paged(
            "a.html", None, None, chain, _quality_ok, lambda img: ("VLM", {}),
            page_splitter=splitter, page_renderer=lambda p, w: None, page_chain=page_chain)
        assert chain_calls == ["a.html"]
        assert page_chain_calls == []


class TestRealSplitters:
    def test_split_pdf_pages(self, tmp_path):
        import pymupdf
        pdf = tmp_path / "doc.pdf"
        doc = pymupdf.open()
        for _ in range(2):
            page = doc.new_page()
            page.insert_text((72, 72), "Hello")
        doc.save(str(pdf))
        doc.close()
        work = tmp_path / "work"
        work.mkdir()
        pages = pp.split_pdf_pages(str(pdf), str(work))
        assert len(pages) == 2
        assert all(Path(p).exists() for p in pages)

    def test_render_pdf_page(self, tmp_path):
        import pymupdf
        pdf = tmp_path / "page.pdf"
        doc = pymupdf.open()
        doc.new_page().insert_text((72, 72), "Hello")
        doc.save(str(pdf))
        doc.close()
        work = tmp_path / "work"
        work.mkdir()
        image = pp.render_pdf_page(str(pdf), str(work))
        assert image is not None and Path(image).exists()

    def test_default_page_splitter_by_suffix(self, tmp_path):
        assert pp.default_page_splitter("a.html", str(tmp_path)) is None
        assert pp.default_page_splitter("a.docx", str(tmp_path)) is None
        assert pp.default_page_renderer("a.txt", str(tmp_path)) is None

    def test_split_pptx_slides(self, tmp_path):
        from pptx import Presentation
        pptx_path = tmp_path / "deck.pptx"
        prs = Presentation()
        for _ in range(2):
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            slide.shapes.title.text = "标题"
        prs.save(str(pptx_path))
        work = tmp_path / "work"
        work.mkdir()
        slides = pp.split_pptx_slides(str(pptx_path), str(work))
        assert len(slides) == 2
        assert all(Path(s).exists() for s in slides)
